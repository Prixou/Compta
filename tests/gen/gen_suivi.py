# Tableau de suivi TVA fictif (même forme qu'un classeur de cabinet) pour les tests d'import et du suivi mensuel.
import datetime, random, sys
import openpyxl
from openpyxl.styles import PatternFill
random.seed(5)
OUT = sys.argv[1] if len(sys.argv) > 1 else 'import-dossiers.xlsx'
NOMS = ['SCI DES TILLEULS', 'SCI LE MAS BLEU', 'SCI DU CANAL', 'SCI LES PINS', 'SCI ROCHEFORT', 'SCI DU MOULIN', 'SCI LES ARCADES', 'SCI PLEIN SUD',
        'SCI LA GARRIGUE', 'SCI DES OLIVIERS', 'SCI CAP OUEST', 'SCI DU PARC', 'SCI LES VIGNES', 'SCI HORIZON', 'SCI LE CLOS',
        'SARL BOULANGERIE MARTIN', 'SARL GARAGE DU CENTRE', 'SARL IMMOBILIERE DES COTEAUX', 'SARL PLOMBERIE DURAND', 'SARL LES DELICES', 'SARL TRANSPORTS RAPIDES', 'SARL ATELIER BOIS',
        'SELARL CABINET DENTAIRE SOLEIL', 'SELARL KINE PLUS', 'SELARL VETO DU PONT', 'SELARL AVOCATS ASSOCIES',
        'SAS TECH SOLUTIONS', 'SAS BIO MARCHE', 'SAS EVENTS SUD', 'SAS DIGITAL WEB', 'SELAS LABO ANALYSES', 'SELAS RADIOLOGIE EST', 'SELAS OPTIQUE CENTRE',
        'EURL CONSEIL PLUS', 'EURL FLEURS ET JARDINS', 'EARL DOMAINE DES SOURCES', 'SMC MUTUELLE LOCALE', 'EI DUPONT ARTISAN']
STAT = ['SCI'] * 15 + ['SARL'] * 7 + ['SELARL'] * 4 + ['SAS'] * 4 + ['SELAS'] * 3 + ['EURL'] * 2 + ['EARL', 'SMC', 'EI']
TRIM = ['SCI LES ARCADES', 'SCI DU MOULIN', 'SCI HORIZON', 'SCI LE CLOS', 'SAS EVENTS SUD', 'EURL CONSEIL PLUS', 'SMC MUTUELLE LOCALE']
TVA = ['CA12' if n == 'EARL DOMAINE DES SOURCES' else 'T' if n in TRIM else 'M' for n in NOMS]
CLOTURE = {'SELARL KINE PLUS': datetime.datetime(2026, 3, 31)}
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Suivi TVA 2026'
ws.append(['N°DOSSIER', 'STATUT', 'DOSSIERS', 'RESPONSABLE (M)', 'COLLABORATEUR (C)', 'SUPERVISEUR (C)', 'SIREN', 'TVA', 'JOUR TVA', 'IS/IR'] + list(range(1, 13)))
ok = PatternFill('solid', fgColor='009DC3E6')
na = PatternFill('darkUp', fgColor='00595959')
for i, nom in enumerate(NOMS):
    t = TVA[i]
    resp = 'AAA' if t == 'M' else 'BBB' if t == 'T' else 'FFF'
    jour = None if t == 'CA12' else (24 if t == 'T' else (16 if i == 5 else 21))
    ws.append([f'{900100 + i * 7:06d}', STAT[i], nom, resp, ['CCC', 'DDD', 'EEE'][i % 3], 'ZZZ', f'9{i:02d}{i * 37 % 1000:03d}{i:03d}00010', t, jour, CLOTURE.get(nom, datetime.datetime(2026, 12, 31))] + [None] * 12)
    row = ws.max_row
    for m in range(1, 13):
        c = ws.cell(row=row, column=10 + m)
        if t == 'T' and m % 3:
            c.fill = na
        elif m <= 8 and (t == 'M' or (t == 'T' and m % 3 == 0)):
            c.value = 'OK'
            c.fill = ok
wb.save(OUT)
print(OUT, len(NOMS), 'dossiers')
